
import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}
MAX_DOCUMENT_SIZE = 10 * 1024 * 1024  # 10 MB


def validate_document(filename, file_size=None):
    """Validate supported document extensions and optional file size."""
    if not isinstance(filename, str) or not filename.strip():
        raise ValueError("A filename is required.")

    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError("Only PDF, DOCX, and TXT files are supported.")

    if file_size is not None:
        if file_size <= 0:
            raise ValueError("The uploaded file is empty.")

        if file_size > MAX_DOCUMENT_SIZE:
            raise ValueError("The file exceeds the 10 MB size limit.")

    return True


def extract_text_from_txt(file_path):
    """Extract UTF-8 text from a TXT file."""
    path = Path(file_path)

    try:
        return path.read_text(encoding="utf-8-sig").strip()
    except UnicodeDecodeError as exc:
        raise ValueError("The TXT file must use UTF-8 encoding.") from exc
    except OSError as exc:
        raise ValueError("Unable to read this TXT file.") from exc


def extract_text_from_pdf(file_path):
    """Extract text from a PDF file."""
    try:
        reader = PdfReader(str(file_path))

        if reader.is_encrypted:
            raise ValueError("Encrypted PDFs are not supported.")

        extracted_pages = []

        for page in reader.pages:
            text = page.extract_text()

            if text:
                extracted_pages.append(text)

        return "\n\n".join(extracted_pages).strip()

    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Unable to read this PDF file.") from exc


def extract_text_from_docx(file_path):
    """Extract paragraphs and table contents from a DOCX file."""
    try:
        document = Document(str(file_path))
        parts = []

        # Extract normal paragraphs.
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                parts.append(text)

        # Extract table cells, which often contain important document data.
        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip().replace("\n", " ")
                    for cell in row.cells
                ]

                if any(cells):
                    parts.append(" | ".join(cells))

        return "\n\n".join(parts).strip()

    except Exception as exc:
        raise ValueError(
            "Unable to read this DOCX file. Make sure it is a valid Word document."
        ) from exc


def extract_document_text(file_path):
    """Validate a document and extract text according to its extension."""
    path = Path(file_path)

    validate_document(path.name)

    if not path.is_file():
        raise ValueError("The document does not exist.")

    file_size = path.stat().st_size

    if file_size <= 0:
        raise ValueError("The uploaded file is empty.")

    if file_size > MAX_DOCUMENT_SIZE:
        raise ValueError("The file exceeds the 10 MB size limit.")

    extension = path.suffix.lower()

    if extension == ".txt":
        text = extract_text_from_txt(path)
    elif extension == ".pdf":
        text = extract_text_from_pdf(path)
    elif extension == ".docx":
        text = extract_text_from_docx(path)
    else:
        # Defensive check in case supported extensions change later.
        raise ValueError("Only PDF, DOCX, and TXT files are supported.")

    if not text.strip():
        raise ValueError(
            "No readable text was found. Scanned PDFs may require OCR."
        )

    return text


def clean_text(text):
    """Normalize whitespace while preserving paragraph boundaries."""
    if not isinstance(text, str):
        raise ValueError("Text must be a string.")

    text = text.replace("\x00", " ")
    text = text.replace("\r\n", "\n").replace("\r", "\n")

    lines = [line.strip() for line in text.split("\n")]
    text = "\n".join(lines)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ \t]+", " ", text)

    return text.strip()


def chunk_text(text, chunk_size=1000, overlap=150):
    """Split text into overlapping chunks without dropping characters."""
    if not isinstance(text, str):
        raise ValueError("Text must be a string.")

    if not isinstance(chunk_size, int) or chunk_size <= 0:
        raise ValueError("chunk_size must be a positive integer.")

    if not isinstance(overlap, int) or overlap < 0:
        raise ValueError("overlap must be a non-negative integer.")

    if overlap >= chunk_size:
        raise ValueError("overlap must be smaller than chunk_size.")

    text = clean_text(text)

    if not text:
        return []

    chunks = []
    start = 0
    text_length = len(text)

    while start < text_length:
        end = min(start + chunk_size, text_length)
        chunk = text[start:end].strip()

        if chunk:
            chunks.append(chunk)

        if end == text_length:
            break

        start = end - overlap

    return chunks
